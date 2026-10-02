"""
workbook.py - the downloadable expense workbook and reading it back when uploaded.
Also reads text out of an uploaded Word (.docx) or .txt file for the description boxes.

Upload accepts the app's own template AND the older RD_Credit_Questionnaire layout
(same sheet names: Wages QRE, Supplies QRE, Cloud_Compute QRE, Contractor US QRE,
Contractor Foreign QRE). Columns are found by their header text.
"""
import io

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from study import CATEGORIES

TEMPLATE_HEADERS = {
    "wages":    ["Employee Name", "Officer (Yes/No)", "Job Title", "State", "Total Taxable Wages", "Qualified %"],
    "supplies": ["Vendor Name", "Type", "State", "Amount Used for R&D"],
    "cloud":    ["Vendor Name", "Type", "State", "Amount Used for R&D"],
    "contract": ["Contractor Name", "Work Performed", "State", "Amount Paid", "Qualified %"],
    "foreign":  ["Contractor Name", "Work Performed", "Country", "Amount Paid", "R&E %"],
}


def template_bytes(title="R&D Feasibility Study"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Instructions"
    lines = [
        f"{title} – expense workbook",
        "",
        "Fill in one row per employee / vendor / contractor on each tab, then upload this file on step 4.",
        "Qualified % and R&E %: enter as a percentage (e.g. 50%). Leave a row blank to skip it.",
        "Wages – Officer (Yes/No): choose Yes for corporate officers (used for Form 6765 line 38). Blank = No.",
        "Supplies and Computer / Cloud: enter only the amount used for R&D (no percentage).",
        "U.S. contract research: enter the full amount paid – the app applies the 65% limit.",
        "Foreign contractors are kept outside the credit (Section 174, 15-year amortization).",
        "Do not rename the tabs or the header row.",
    ]
    for i, t in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t)
    ws["A1"].font = Font(bold=True, size=13, color="003366")
    ws.column_dimensions["A"].width = 110

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="003366")
    for cat, headers in TEMPLATE_HEADERS.items():
        s = wb.create_sheet(CATEGORIES[cat]["sheet"])
        for col, h in enumerate(headers, 1):
            c = s.cell(row=1, column=col, value=h)
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(horizontal="center")
            s.column_dimensions[c.column_letter].width = 30 if col <= 2 else 20
        if cat == "wages":
            dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
            dv.error, dv.errorTitle = "Choose Yes or No", "Officer"
            s.add_data_validation(dv)
            dv.add("B2:B201")
            s.column_dimensions["B"].width = 16
        for r in range(2, 202):
            for col, h in enumerate(headers, 1):
                if "%" in h:
                    s.cell(row=r, column=col).number_format = "0%"
                elif "Amount" in h or "Wages" in h:
                    s.cell(row=r, column=col).number_format = "#,##0.00"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _find(headers, *words, exclude=()):
    for i, h in enumerate(headers):
        hl = h.lower()
        if any(w in hl for w in words) and not any(x in hl for x in exclude):
            return i
    return None


def _pct(v):
    """Excel stores 50% as 0.5; a typed 50 means 50%."""
    if v is None or v == "":
        return ""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(v)
    return f"{f * 100:g}%" if abs(f) <= 1 else f"{f:g}%"


def _amt(v):
    if v is None or v == "":
        return ""
    try:
        return f"{float(v):,.2f}"
    except (TypeError, ValueError):
        return str(v)


def parse_upload(file_storage):
    """Return {cat: [row dicts]} for every category tab found in the workbook."""
    wb = openpyxl.load_workbook(file_storage, data_only=True)
    found = {}
    for cat, cfg in CATEGORIES.items():
        if cfg["sheet"] not in wb.sheetnames:
            continue
        ws = wb[cfg["sheet"]]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h or "").strip() for h in rows[0]]
        i_name = 0
        i_col2 = _find(headers, "job title", "work performed", "type")
        i_off = _find(headers, "officer") if cfg.get("officer") else None
        i_state = _find(headers, "state") if cat != "foreign" else _find(headers, "country")
        if i_state is None:
            i_state = _find(headers, "country", "state")
        i_amt = _find(headers, "amount", "taxable wages", "wages", exclude=("calculation", "qre", "total wages"))
        i_pct = _find(headers, "%")
        out = []
        for r in rows[1:]:
            def cell(i):
                return r[i] if i is not None and i < len(r) else None
            name, amount = cell(i_name), cell(i_amt)
            if (name in (None, "")) and (amount in (None, "")):
                continue
            row = {"name": str(name or "").strip(), "col2": str(cell(i_col2) or "").strip(),
                   "state": str(cell(i_state) or "").strip(), "amount": _amt(amount), "pct": ""}
            if cfg.get("officer"):
                row["officer"] = "Yes" if str(cell(i_off) or "").strip().lower() in ("yes", "y", "true", "x") else "No"
            if cfg["pct"]:
                row["pct"] = _pct(cell(i_pct))
            elif i_pct is not None and cell(i_pct) not in (None, ""):
                # older questionnaire: supplies/cloud had a Qualified % - convert to amount used for R&D
                try:
                    p = float(cell(i_pct))
                    p = p if p <= 1 else p / 100
                    row["amount"] = _amt(float(amount or 0) * p)
                except (TypeError, ValueError):
                    pass
            out.append(row)
        found[cat] = out
    return found


def read_text(file_storage):
    """Text from an uploaded .txt or .docx file ('' if unreadable)."""
    name = (file_storage.filename or "").lower()
    data = file_storage.read()
    if not data:
        return ""
    if name.endswith(".docx"):
        try:
            import docx
            d = docx.Document(io.BytesIO(data))
            return "\n".join(p.text for p in d.paragraphs).strip()
        except Exception:
            return ""
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return ""
